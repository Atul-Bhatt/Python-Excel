from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.drawing.image import Image
from openpyxl import load_workbook
# from PIL import Image
wb = load_workbook('Pie.xlsx')
ws = wb.active

tab = Table(displayName='Table', ref='A1:B5')
style = TableStyleInfo(name='TableStyleMedium9', showFirstColumn=False, showLastColumn=False,
                       showRowStripes=True, showColumnStripes=True)
tab.tableStyleInfo = style
ws.add_table(tab)

img = Image('madecraft.jpg')
img.height = img.height * .25
img.width = img.width * .25
ws.add_image(img, 'C1')

wb.save('new_image.xlsx')
